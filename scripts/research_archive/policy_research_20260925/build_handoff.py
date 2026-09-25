import hashlib
import json
import shutil
from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path.cwd()
WORK = ROOT/'tmp/policy_research_20260925'
OUT = ROOT/'outputs/policy_review_20260925'
OUT.mkdir(parents=True, exist_ok=True)
EVIDENCE = OUT/'evidence'
EVIDENCE.mkdir(exist_ok=True)
DATE = '2026-09-25'
REGISTRY = 'https://www.daejeon.go.kr/drh/board/boardNormalView.do?boardId=normal_0096&menuSeq=1631&ntatcSeq=1521961064'
PARK = 'https://www.daejeon.go.kr/ani/AniContentsHtmlView.do?menuSeq=7367'
FACILITIES = 'https://www.daejeon.go.kr/ani/AniContentsHtmlView.do?menuSeq=7368'
EDUCATION = 'https://www.daejeon.go.kr/ani/AniContentsHtmlView.do?menuSeq=7369'
AIRBNB = 'https://www.airbnb.co.kr/rooms/1256582248927052245'


def save(name, data):
    (OUT/name).write_text(json.dumps(data, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


# Original bytes are copied, not rewritten as new spreadsheets.
evidence = []
for name in ['restaurant_latest','restaurant_xlsx','park_policy','park_facilities','education','chuseok','chuseok_notice','agility_closure']:
    meta=json.loads((WORK/'sources'/f'{name}.json').read_text(encoding='utf-8'))
    src=ROOT/meta['file']
    suffix='.xlsx' if name=='restaurant_xlsx' else '.jpg' if name=='chuseok_notice' else src.suffix
    dst=EVIDENCE/(name+suffix)
    shutil.copyfile(src,dst)
    evidence.append({**{k:v for k,v in meta.items() if k not in ['links','file']},'file':dst.relative_to(OUT).as_posix()})
for src in sorted((WORK/'tourapi_details').glob('*.json')):
    if src.name.endswith('_review.json'):
        continue
    dst=EVIDENCE/src.name
    shutil.copyfile(src,dst)
    evidence.append(dict(file=dst.relative_to(OUT).as_posix(),sha256=hashlib.sha256(dst.read_bytes()).hexdigest(),
                         checked_at=DATE,source='한국관광공사 KorPetTourService2',
                         request='contentId='+src.stem.split('_')[0],operation=src.stem.split('_',1)[1] if '_' in src.stem else 'detailPetTour2'))
evidence.append(dict(url=AIRBNB,checked_at=DATE,access='web tool readable host listing; no local HTML archive',
                     excerpt='반려동물 동반 허용',summary='호스트 빨간대문의 독채 상품. 반려견 용품과 체크인 15시·체크아웃 11시 안내. 체중·마릿수 숫자는 확인되지 않음.'))
save('sources.json',evidence)

wb=openpyxl.load_workbook((EVIDENCE/'restaurant_xlsx.xlsx'),data_only=True)
ws=wb.worksheets[0]
merged={}
for region in ws.merged_cells.ranges:
    if region.min_row<4:
        continue
    for row in range(region.min_row,region.max_row+1):
        for col in range(region.min_col,region.max_col+1):
            merged[(row,col)]=(ws.cell(region.min_row,region.min_col).value,ws.cell(region.min_row,region.min_col).coordinate)
registrations=[]
restaurant_venues={}
for row in range(4,73):
    vals=[ws.cell(row,col).value for col in range(1,7)]
    original=list(vals)
    fills=[]
    for col in [4,5,6]:
        if (row,col) in merged and vals[col-1] is None:
            vals[col-1],anchor=merged[(row,col)]
            fills.append(f'{ws.cell(row,col).coordinate} <- {anchor} (병합 셀)')
    number,authority,kind,name,address,note=vals
    assert number and name and address
    key=(name,address)
    venue_id='food_'+hashlib.sha256((name+'|'+address).encode()).hexdigest()[:12]
    authority_resolved=authority or address.split()[1]
    registration=dict(registration_id=f'dj_20260907_{number:03d}',venue_id=venue_id,
                      source_row=row,source_locator=f'대전!A{row}:F{row}',raw_values=original,
                      name=name,address=address,business_type=kind,authority_raw=authority,
                      authority_resolved=authority_resolved,
                      transform_notes=fills+([] if authority else ['원본 관할기관 공란; 주소의 구 이름으로 보완, 원본 공란 보존']),
                      registry_as_of='2026-09-07',checked_at=DATE,source_url=REGISTRY,source_note=note)
    registrations.append(registration)
    if key not in restaurant_venues:
        restaurant_venues[key]=dict(venue_id=venue_id,name=name,address=address,category='restaurant',
             latitude=None,longitude=None,registration_ids=[],business_types=[],checked_at=DATE,
             pet_allowed=True,policy_status='partial',max_dogs=None,max_weight_kg=None,
             dogs_unlimited=False,weight_unlimited=False,allowed_space=None,serves_meals=None,
             source_url=REGISTRY,active=False,
             pending=['정확한 좌표','식사 메뉴 제공 여부','실내/테라스 등 동반 공간','마릿수·체중 및 견종 제한','영업시간·휴무·예약','영업 지속 여부'])
    item=restaurant_venues[key]
    item['registration_ids'].append(registration['registration_id'])
    item['business_types'].append(kind)
save('restaurant_registrations.json',registrations)
save('restaurant_candidates.json',list(restaurant_venues.values()))


def tour(content_id,operation):
    return json.loads((WORK/'tourapi_details'/f'{content_id}_{operation}_review.json').read_text(encoding='utf-8'))['records'][0]


red=tour('3533154','detailCommon2')
park=tour('2930677','detailCommon2')
review=[]
imports=[]


def add(id,venue_id,base,category,name,scope,url,summary,pending,requirements=None,mode='unknown',recurring=False,allowed=True):
    review.append(dict(id=id,venue_id=venue_id,name=base['title'],offering_name=name,scope_type=scope,
                       category=category,source_url=url,checked_at=DATE,confirmed_summary=summary,
                       review_status='partial',pet_allowed=allowed,pending=pending,active=False))
    imports.append(dict(id=id,venue_id=venue_id,name=base['title'],address=base['addr1'],
                        latitude=float(base['mapy']),longitude=float(base['mapx']),category=category,
                        description=summary+' 미확인: '+'; '.join(pending),active=False,product_name=name,
                        participation_mode=mode,recurring=recurring,serves_meals=False,
                        official_url=url,schedule_note='방문일 운영·휴무·예약과 프로그램 회차는 재확인 필요. 동봉 검토본 참조.',
                        policy=dict(verified=False,pet_allowed=allowed is True,max_dogs=None,max_weight_kg=None,
                                    dogs_unlimited=False,weight_unlimited=False,requirements=requirements or [],
                                    source_url=url,source_quote=summary,checked_at=DATE)))

add('tour_3533154','tour_3533154',red,'lodging','빨간대문 독채','room',AIRBNB,
    '호스트 예약 페이지와 TourAPI에서 반려견 동반 확인. 독채 상품, 15시 입실·11시 퇴실.',
    ['반려견 마릿수 상한','체중 상한 및 경계값','추가요금·견종 제한','방문일 예약 가능 여부'])
common=['동물등록·광견병 예방접종 확인','생후 2개월 이상','제한 견종 및 공격성 개체 입장 제한','놀이터 밖 목줄 2m 이내','실내 매너밴드']
for key,name,rule in [
    ('small_outdoor','중·소형견 야외놀이터','체고 40cm 미만'),
    ('large_outdoor','대형견 야외놀이터','체고 40cm 이상'),
    ('indoor','실내놀이터','체고 40cm 미만; 실내 매너밴드'),
    ('agility_park','어질리티 파크','대형견만 입장; 체고 40cm 이상')]:
    add('park_'+key,'tour_2930677',park,'activity',name,'space',FACILITIES,
        rule+'인 반려견의 공간 이용을 공식 시설 안내에서 확인.',
        ['보호자/일행당 허용 마릿수','체중 제한 유무','체고·견종·등록/접종을 현재 5개 입력으로 판별 불가','방문일 운영 여부'],
        common+[rule],mode='dog_participates')
for key,name,mode in [
    ('behavior','행동교정교육','dog_participates'),
    ('agility_class','반려견 스포츠 체험교육','dog_participates'),
    ('snack','수제간식 만들기','unknown'),
    ('socialization','반려견 통합사회화','unknown')]:
    add('program_'+key,'tour_2930677',park,'activity',name,'program',EDUCATION,
        '2026년 하반기 9~11월 운영 안내에 포함. OK예약 선착순 접수, 대전 시민 우선.',
        ['현재 모집 회차·시간·잔여석','마릿수·체중/체고 조건','외지인 신청 자격','실제 반려견 참여/동반 범위','소요시간'],
        ['회차별 사전예약','대전 시민 우선'],mode=mode,allowed=True if mode=='dog_participates' else None)
save('policy_review.json',review)
save('places_import_inactive.json',imports)

inventory=json.loads((WORK/'inventory.json').read_text(encoding='utf-8'))
root_csv=ROOT/'region_monthly.csv'
inventory['files'].append(dict(path='region_monthly.csv',bytes=root_csv.stat().st_size,
    sha256=hashlib.sha256(root_csv.read_bytes()).hexdigest(),note='검증본과 같은 SHA-256인지 대조; 중복 적재 금지'))
inventory['snapshot_date']=DATE
inventory['policy_review_files']='outputs/policy_review_20260925/sources.json'
save('workspace_data_inventory.json',inventory)

stats=dict(registrations=len(registrations),unique_restaurant_name_address_pairs=len(restaurant_venues),
           business_type_counts=dict(Counter(x['business_type'] for x in registrations)),
           district_registration_counts=dict(Counter(x['authority_resolved'] for x in registrations)),
           merged_source_rows=[x['source_row'] for x in registrations if any('병합' in t for t in x['transform_notes'])],
           missing_authority_rows=[x['source_row'] for x in registrations if not x['authority_raw']],
           reviewed_offerings=len(review),inactive_import_records=len(imports),active_import_records=0)
save('counts.json',stats)
print(json.dumps(stats,ensure_ascii=False,indent=2))
