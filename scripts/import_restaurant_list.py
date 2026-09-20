import argparse
import csv
import hashlib
from pathlib import Path
from scripts.import_manual_places import write_new


def read_rows(path, sheet, header_row):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for _ in range(header_row - 1):
                next(reader)
            headers = [str(v).strip() for v in next(reader)]
            yield from (dict(zip(headers, row)) for row in reader)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            rows = ws.iter_rows(min_row=header_row, values_only=True)
            headers = [str(v or "").strip() for v in next(rows)]
            yield from (dict(zip(headers, row)) for row in rows)
        finally:
            wb.close()


def main():
    parser = argparse.ArgumentParser(description="대전시 공개 명부의 실제 헤더를 지정해 미검토 후보 JSON으로 변환")
    parser.add_argument("input", type=Path)
    parser.add_argument("--name-column", required=True)
    parser.add_argument("--address-column", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--source-url", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    records = {}
    for row in read_rows(args.input, args.sheet, args.header_row):
        if args.name_column not in row or args.address_column not in row:
            parser.error("지정한 이름·주소 헤더를 찾지 못했습니다.")
        name, address = str(row[args.name_column] or "").strip(), str(row[args.address_column] or "").strip()
        if not name or not address.startswith(("대전 ", "대전광역시 ")):
            continue
        key = "city_" + hashlib.sha256((name + "|" + address).encode()).hexdigest()[:16]
        records[key] = {"id":key, "venue_id":key, "name":name, "address":address,
                        "category":"restaurant", "latitude":None, "longitude":None, "active":False,
                        "serves_meals":False, "description":"시 공개 명부 후보. 좌표·식사 메뉴·규정 검토 필요.",
                        "policy":{"verified":False, "pet_allowed":False, "source_url":args.source_url,
                                  "source_quote":"등록 명부 확인. 상세 동반 규정은 미검토."}}
    write_new(args.output, list(records.values()))
    print(f"미검토 후보 {len(records)}개 생성. 좌표·메뉴·마릿수·체중 근거를 보완한 뒤 적재하세요.")


if __name__ == "__main__":
    main()
